"""
초안 생성기 - LLM을 활용해 공고 분석 결과를 DOCX/MD 파일로 자동 생성한다.
"""
import logging
import os
import re
from datetime import datetime
from pathlib import Path
from typing import List, Optional

from ..crawler.base_crawler import Announcement


logger = logging.getLogger(__name__)


class DraftGenerator:
    """
    공고 분석 초안 파일 생성기.
    - 지원사업 요약 브리핑 (1페이지)
    - 사업계획서 초안 (목적/필요성/추진계획/예산안)
    - 체크리스트 (필수 서류, 자격 조건)
    """

    def __init__(self, config: dict, openai_config: Optional[dict] = None,
                 company_profile: Optional[dict] = None):
        self.output_dir = Path(config.get("output_dir", "output/drafts"))
        self.output_dir.mkdir(parents=True, exist_ok=True)
        self.output_format = config.get("output_format", "docx")
        self.include_summary = config.get("include_summary", True)
        self.include_plan = config.get("include_business_plan", True)
        self.include_checklist = config.get("include_checklist", True)
        self.openai_config = openai_config or {}
        self.company_profile = company_profile or {}
        self._openai_client = None

    def generate(self, announcement: Announcement,
                 downloaded_files: Optional[List[str]] = None) -> Optional[str]:
        """
        공고에 대한 초안 파일을 생성한다.

        Args:
            announcement: 분석할 공고
            downloaded_files: 다운로드된 첨부파일 경로 목록

        Returns:
            생성된 초안 파일 경로 (실패 시 None)
        """
        logger.info("초안 생성 시작: %s", announcement.title)
        try:
            content = self._build_draft_content(announcement, downloaded_files or [])
            return self._save_draft(announcement, content)
        except Exception as exc:
            logger.error("초안 생성 실패 (%s): %s", announcement.title, exc)
            return None

    # ------------------------------------------------------------------ #
    # 초안 내용 구성                                                         #
    # ------------------------------------------------------------------ #

    def _build_draft_content(self, ann: Announcement, files: List[str]) -> dict:
        """LLM으로 초안 섹션을 생성한다."""
        sections = {}

        if self.include_summary:
            sections["summary"] = self._generate_section(
                ann, "요약 브리핑",
                self._summary_prompt(ann),
            )

        if self.include_plan:
            sections["business_plan"] = self._generate_section(
                ann, "사업계획서 초안",
                self._business_plan_prompt(ann),
            )

        if self.include_checklist:
            sections["checklist"] = self._generate_section(
                ann, "체크리스트",
                self._checklist_prompt(ann),
            )

        return sections

    def _generate_section(self, ann: Announcement, section_name: str, prompt: str) -> str:
        """GPT를 호출하여 섹션 내용을 생성한다."""
        try:
            client = self._get_openai_client()
            model = self.openai_config.get("model", "gpt-4o")
            response = client.chat.completions.create(
                model=model,
                messages=[{"role": "user", "content": prompt}],
                max_tokens=self.openai_config.get("max_tokens", 2000),
                temperature=0.4,
            )
            return response.choices[0].message.content.strip()
        except Exception as exc:
            logger.warning("[%s] %s 생성 실패 (LLM 미사용): %s", ann.title, section_name, exc)
            return self._fallback_section(ann, section_name)

    def _fallback_section(self, ann: Announcement, section_name: str) -> str:
        """LLM 없이 기본 템플릿으로 섹션 생성"""
        if section_name == "요약 브리핑":
            return (
                f"■ 공고명: {ann.title}\n"
                f"■ 주관기관: {ann.organization}\n"
                f"■ 지원분야: {ann.category}\n"
                f"■ 지원대상: {ann.target}\n"
                f"■ 지원금액: {ann.support_amount}\n"
                f"■ 접수기간: {ann.start_date} ~ {ann.end_date}\n"
                f"■ 공고URL: {ann.url}\n"
                f"■ 적합도: {ann.eligibility_score:.1f}점\n"
                f"■ 판단근거: {ann.eligibility_reason}\n"
            )
        elif section_name == "체크리스트":
            return (
                "□ 공고문 전문 확인\n"
                "□ 지원 자격 요건 충족 여부 확인\n"
                "□ 사업자등록증 준비\n"
                "□ 법인등기부등본 준비\n"
                "□ 재무제표 (최근 1~2년) 준비\n"
                "□ 관련 인증서 사본 준비\n"
                "□ 신청서 양식 다운로드\n"
                "□ 사업계획서 작성\n"
                "□ 온라인 신청 또는 방문 접수\n"
            )
        return f"[{section_name}] 공고를 확인하고 내용을 직접 작성하세요.\n공고 URL: {ann.url}"

    # ------------------------------------------------------------------ #
    # LLM 프롬프트 구성                                                     #
    # ------------------------------------------------------------------ #

    def _summary_prompt(self, ann: Announcement) -> str:
        company_name = self.company_profile.get("name", "우리 기업")
        return (
            "당신은 정부 지원사업 전문 컨설턴트입니다.\n"
            f"아래 지원사업 공고를 분석하여 {company_name} 담당자가 즉시 이해할 수 있는 "
            "1페이지 요약 브리핑을 한국어로 작성하세요.\n"
            "포함 항목: 사업 목적, 주요 지원 내용, 지원 금액, 신청 자격, 접수 방법, 주요 일정\n\n"
            f"공고 제목: {ann.title}\n"
            f"주관 기관: {ann.organization}\n"
            f"지원 대상: {ann.target}\n"
            f"지원 금액: {ann.support_amount}\n"
            f"공고 내용:\n{ann.raw_content[:3000]}"
        )

    def _business_plan_prompt(self, ann: Announcement) -> str:
        profile = self.company_profile
        return (
            "당신은 정부 지원사업 전문 컨설턴트입니다.\n"
            "아래 지원사업 공고와 기업 정보를 바탕으로 사업계획서 초안을 한국어로 작성하세요.\n"
            "구성: 1.사업 목적 및 필요성, 2.추진 계획(월별 세부 일정), "
            "3.기대 효과, 4.예산 계획(개요)\n\n"
            f"[기업 정보]\n"
            f"업종: {profile.get('industry', '')} / {profile.get('sub_industry', '')}\n"
            f"설립연도: {profile.get('established_year', '')}\n"
            f"직원 수: {profile.get('employee_count', '')}명\n"
            f"보유 인증: {', '.join(profile.get('certifications', []))}\n\n"
            f"[공고 정보]\n"
            f"공고명: {ann.title}\n"
            f"지원 내용: {ann.description}\n"
            f"공고 전문:\n{ann.raw_content[:3000]}"
        )

    def _checklist_prompt(self, ann: Announcement) -> str:
        profile = self.company_profile
        return (
            "당신은 정부 지원사업 전문 컨설턴트입니다.\n"
            "아래 공고를 분석하여 신청을 위한 체크리스트를 한국어로 작성하세요.\n"
            "포함 항목: ① 필수 제출 서류 목록, ② 자격 조건 충족 여부 체크, "
            "③ 신청 전 확인 사항, ④ 주의 사항\n\n"
            f"[기업 현황]\n"
            f"업종: {profile.get('industry', '')}\n"
            f"직원: {profile.get('employee_count', '')}명\n"
            f"인증: {', '.join(profile.get('certifications', []))}\n\n"
            f"[공고 정보]\n"
            f"공고명: {ann.title}\n"
            f"지원 대상: {ann.target}\n"
            f"공고 전문:\n{ann.raw_content[:3000]}"
        )

    # ------------------------------------------------------------------ #
    # 파일 저장                                                             #
    # ------------------------------------------------------------------ #

    def _save_draft(self, ann: Announcement, sections: dict) -> str:
        date_str = datetime.now().strftime("%Y%m%d")
        safe_title = re.sub(r'[\\/:*?"<>|\s]', "_", ann.title)[:50]
        filename = f"{ann.site_id}_{safe_title}_{date_str}"

        if self.output_format == "docx":
            return self._save_docx(filename, ann, sections)
        else:
            return self._save_markdown(filename, ann, sections)

    def _save_markdown(self, filename: str, ann: Announcement, sections: dict) -> str:
        path = self.output_dir / f"{filename}.md"
        lines = [
            f"# {ann.title}",
            f"",
            f"- **사이트**: {ann.site_name}",
            f"- **주관기관**: {ann.organization}",
            f"- **적합도**: {ann.eligibility_score:.1f}점",
            f"- **공고 URL**: {ann.url}",
            f"- **생성일**: {datetime.now().strftime('%Y-%m-%d %H:%M')}",
            "",
        ]

        if "summary" in sections:
            lines += ["---", "## 📋 요약 브리핑", "", sections["summary"], ""]
        if "business_plan" in sections:
            lines += ["---", "## 📄 사업계획서 초안", "", sections["business_plan"], ""]
        if "checklist" in sections:
            lines += ["---", "## ✅ 체크리스트", "", sections["checklist"], ""]

        with open(path, "w", encoding="utf-8") as fp:
            fp.write("\n".join(lines))

        logger.info("MD 초안 저장: %s", path)
        return str(path)

    def _save_docx(self, filename: str, ann: Announcement, sections: dict) -> str:
        path = self.output_dir / f"{filename}.docx"
        try:
            from docx import Document
            from docx.shared import Pt, RGBColor
            from docx.enum.text import WD_ALIGN_PARAGRAPH

            doc = Document()

            # 제목
            title_para = doc.add_heading(ann.title, level=0)
            title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

            # 메타 정보
            doc.add_paragraph(f"사이트: {ann.site_name}  |  주관기관: {ann.organization}")
            doc.add_paragraph(f"적합도: {ann.eligibility_score:.1f}점  |  공고 URL: {ann.url}")
            doc.add_paragraph(f"생성일: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
            doc.add_paragraph("")

            if "summary" in sections:
                doc.add_heading("📋 요약 브리핑", level=1)
                doc.add_paragraph(sections["summary"])

            if "business_plan" in sections:
                doc.add_heading("📄 사업계획서 초안", level=1)
                doc.add_paragraph(sections["business_plan"])

            if "checklist" in sections:
                doc.add_heading("✅ 체크리스트", level=1)
                doc.add_paragraph(sections["checklist"])

            doc.save(str(path))
            logger.info("DOCX 초안 저장: %s", path)

        except ImportError:
            logger.warning("python-docx 미설치, Markdown으로 대체 저장")
            return self._save_markdown(filename, ann, sections)

        return str(path)

    def _get_openai_client(self):
        if self._openai_client is None:
            from openai import OpenAI
            api_key = (
                self.openai_config.get("api_key")
                or os.environ.get("OPENAI_API_KEY")
            )
            if not api_key:
                raise ValueError("OPENAI_API_KEY가 설정되지 않았습니다.")
            self._openai_client = OpenAI(api_key=api_key)
        return self._openai_client
